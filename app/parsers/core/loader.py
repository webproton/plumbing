# parsers/core/loader.py
"""Загрузка файла накладной в универсальный вид: список строк (каждая — список ячеек)."""

import logging
import os

import pandas as pd
from defusedxml import ElementTree as ET  # защита от XXE / «billion laughs» в загружаемых XML
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

YML_PRICE_LIST_MARK = "XML розвантаження B2B"


def load_rows(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext in (".xml", ".yml"):
        try:
            root = ET.parse(file_path).getroot()
        except Exception as exc:
            raise ValueError(f"Не удалось разобрать XML: {exc}") from exc
        if root.tag == "yml_catalog":
            return _yml_catalog_rows(root)
        return _spreadsheet_xml_rows(root)

    if ext == ".xlsx":
        try:
            wb = load_workbook(file_path, data_only=True, read_only=True)
            return [list(r) for r in wb.active.iter_rows(values_only=True)]
        except Exception:
            logger.info("openpyxl не прочитал %s, пробуем pandas", file_path)

    for engine in ("openpyxl", "calamine"):
        try:
            df = pd.read_excel(file_path, header=None, engine=engine)
            return df.fillna("").values.tolist()
        except Exception:
            continue
    raise ValueError("Неподдерживаемый формат Excel или файл повреждён")


def _spreadsheet_xml_rows(root):
    """Старые выгрузки «XML Spreadsheet 2003»: <Row><Cell><Data>…"""
    rows = []
    for row in root.iter():
        if not str(row.tag).endswith("Row"):
            continue
        cells = [(cell.text or "".join(cell.itertext())).strip() for cell in row]
        if any(cells):
            rows.append(cells)
    return rows


def _text(node, tag, default=""):
    child = node.find(tag)
    return child.text.strip() if child is not None and child.text else default


def _yml_catalog_rows(root):
    """
    YML-каталог (прайс поставщика, напр. «Секвоя») → виртуальная таблица.
    Это прайс-лист, а не накладная: количества в нём нет, поэтому колонка «Кількість» пустая —
    при проведении такие строки не приходуются (попадут в отчёт как «нет количества»),
    но цены и артикулы доступны для сопоставления.
    """
    shop = root.find("shop")
    company = (_text(shop, "company") or _text(shop, "name")) if shop is not None else ""

    rows = [
        ["Постачальник / Компанія:", company],
        ["Дата документа:", root.get("date", "")],
        ["Тип файлу:", YML_PRICE_LIST_MARK],
        [],
    ]

    categories = {}
    if shop is not None and shop.find("categories") is not None:
        for cat in shop.find("categories").findall("category"):
            if cat.get("id"):
                categories[cat.get("id")] = cat.text

    rows.append(["Номенклатура (Товар)", "Артикул", "Кількість", "Ціна", "Од. вим.", "Категорія"])

    offers = shop.find("offers") if shop is not None else root.find(".//offers")
    if offers is not None:
        for offer in offers.findall("offer"):
            sku = _text(offer, "vendorCode") or offer.get("id") or ""
            price = _text(offer, "price", "0")
            currency = _text(offer, "currencyId", "UAH")
            rows.append(
                [
                    _text(offer, "name"),
                    sku,
                    "",
                    f"{price} {currency}",
                    "шт",
                    categories.get(_text(offer, "categoryId"), ""),
                ]
            )

    logger.debug("YML: %s строк", len(rows))
    return rows
