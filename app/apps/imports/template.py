"""
Шаблон загрузки товаров (Excel/CSV): скачать → заполнить → загрузить → предпросмотр → применить.

Строка находит существующий товар по артикулу, затем по штрихкоду — и обновляет его,
иначе создаёт новый. Пустая ячейка у существующего товара = «не менять».
Колонка «Остаток» — фактический остаток на выбранном складе: проводится инвентаризацией
(движение в журнале склада), поэтому повторная загрузка того же файла остатки не удваивает.
"""

import csv
import io
import re
from decimal import Decimal

from django.db import IntegrityError, transaction
from django.db.models.functions import Lower
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from app.apps.products.models import Brand, Category, Product, ProductVariant
from app.apps.warehouse import stock
from app.apps.warehouse.models import ProductBarcode, Warehouse

from .fields import to_bool, to_decimal, to_text, to_urls
from .formats import FormatError, detect_format, read_table
from .services import discount_fields, editable_variants

MAX_ROWS = 20_000
MAX_EXPORT = 20_000

COLUMNS = [
    (
        "sku",
        gettext_lazy("Артикул"),
        gettext_lazy("Уникальный код товара. Если товар с таким артикулом есть — он обновится"),
    ),
    ("barcode", gettext_lazy("Штрихкод"), gettext_lazy("EAN-13 или другой код со штрихкода")),
    ("name", gettext_lazy("Название"), gettext_lazy("Обязательно для новых товаров")),
    ("variant", gettext_lazy("Вариант"), gettext_lazy("Цвет, размер и т. п., если у товара несколько вариантов")),
    (
        "group",
        gettext_lazy("Код группы"),
        gettext_lazy("Строки с одинаковым кодом группы станут вариантами одного товара"),
    ),
    ("category", gettext_lazy("Категория"), gettext_lazy("Категория на сайте; если такой нет — будет создана")),
    ("brand", gettext_lazy("Бренд"), ""),
    ("price", gettext_lazy("Цена"), gettext_lazy("Цена продажи, грн. Обязательно для новых товаров")),
    ("old_price", gettext_lazy("Старая цена"), gettext_lazy("Цена до скидки — на сайте будет зачёркнута")),
    ("cost", gettext_lazy("Закупочная цена"), gettext_lazy("Себестоимость, грн")),
    (
        "quantity",
        gettext_lazy("Остаток"),
        gettext_lazy("Фактический остаток на складе, выбранном при загрузке. Пусто — не менять"),
    ),
    ("description", gettext_lazy("Описание"), ""),
    (
        "photos",
        gettext_lazy("Фото"),
        gettext_lazy("Ссылки на фото через запятую (загрузятся, если у товара ещё нет фото)"),
    ),
    ("visible", gettext_lazy("Показывать на сайте"), gettext_lazy("да / нет. Пусто — по настройке загрузки")),
]
HEADERS = {field: header for field, header, _hint in COLUMNS}
_BARCODE = re.compile(r"^[0-9A-Za-z.\-]{1,50}$")


# ---------------------------------------------------------------------------
# Файл шаблона
# ---------------------------------------------------------------------------


def _product_rows(seller):
    qs = (
        editable_variants(seller).select_related("product__category", "brand").order_by("product_id", "id")[:MAX_EXPORT]
    )
    for v in qs:
        p = v.product
        price, old = v.price, v.original_price if v.discount else None
        yield {
            "sku": v.code or "",
            "barcode": v.barcode or "",
            "name": p.name or "",
            "variant": (v.variant_name or "") if v.variant_name != p.name else "",
            "group": f"P{p.pk}",
            "category": p.category.name if p.category else "",
            "brand": v.brand.name if v.brand and v.brand.name else "",
            "price": float(price) if price is not None else "",
            "old_price": float(old) if old else "",
            "cost": float(v.cost_price) if v.cost_price is not None else "",
            "quantity": "",
            "description": p.description or "",
            "photos": "",
            "visible": _("да") if v.available else _("нет"),
        }


def build_template(fmt: str = "xlsx", seller=None) -> tuple[bytes, str]:
    """→ (содержимое, content-type). seller — заполнить строками его товаров (для массового редактирования)."""
    rows = list(_product_rows(seller)) if seller else []
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow([str(h) for h in HEADERS.values()])
        for row in rows:
            writer.writerow([row[f] for f in HEADERS])
        return buffer.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"

    wb = Workbook()
    ws = wb.active
    ws.title = str(_("Товары"))
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8F0FE")
    for col, (field, header, hint) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=str(header))
        cell.font, cell.fill = bold, fill
        if hint:
            cell.comment = Comment(str(hint), "")
        width = {"name": 45, "description": 45, "photos": 40, "category": 22}.get(field, 16)
        ws.column_dimensions[cell.column_letter].width = width
    for r, row in enumerate(rows, start=2):
        for col, field in enumerate(HEADERS, start=1):
            ws.cell(row=r, column=col, value=row[field])
    ws.freeze_panes = "A2"

    info = wb.create_sheet(str(_("Инструкция")))
    info.append([str(_("Колонка")), str(_("Что указывать"))])
    info["A1"].font = info["B1"].font = bold
    for _field, header, hint in COLUMNS:
        info.append([str(header), str(hint)])
    info.append([])
    info.append([str(_("Товар ищется по артикулу, затем по штрихкоду. Найден — обновится, нет — создастся новый."))])
    info.append([str(_("Пустая ячейка у существующего товара — значение не меняется."))])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Разбор и проверка
# ---------------------------------------------------------------------------


def _number(values, field, errors, *, allow_zero=True):
    raw = values.get(field)
    if raw is None or to_text(raw) == "":
        return None
    number = to_decimal(raw)
    if number is None or number < 0 or (not allow_zero and number == 0):
        errors.append(
            _("Некорректное значение в колонке «{column}»: {value}").format(
                column=HEADERS[field], value=to_text(raw, 40)
            )
        )
        return None
    return str(number)


def normalize(values: dict) -> tuple[dict, list, list]:
    errors, warnings = [], []
    data = {
        "sku": to_text(values.get("sku")),
        "barcode": re.sub(r"\s+", "", to_text(values.get("barcode"))),
        "name": to_text(values.get("name")),
        "variant": to_text(values.get("variant"), 255),
        "group": to_text(values.get("group"), 100),
        "category": to_text(values.get("category"), 255),
        "brand": to_text(values.get("brand"), 255),
        "description": str(values.get("description") or "").strip(),
        "photos": to_urls(values.get("photos"))[:10],
        "visible": to_bool(values.get("visible")),
        "price": _number(values, "price", errors, allow_zero=False),
        "old_price": _number(values, "old_price", errors),
        "cost": _number(values, "cost", errors),
        "quantity": _number(values, "quantity", errors),
    }
    if len(data["sku"]) > 50:
        errors.append(_("Артикул длиннее 50 символов"))
    if data["barcode"] and not _BARCODE.match(data["barcode"]):
        errors.append(_("Штрихкод может содержать только цифры, латиницу и дефис (до 50 символов)"))
    if len(data["name"]) > 255:
        errors.append(_("Название длиннее 255 символов"))
    if to_text(values.get("visible")) and data["visible"] is None:
        warnings.append(
            _("«{column}»: не понятно значение, будет применена настройка загрузки").format(column=HEADERS["visible"])
        )
    if to_text(values.get("photos")) and not data["photos"]:
        warnings.append(_("В колонке «Фото» нет ссылок http(s)"))
    return data, errors, warnings


def read_rows(content: bytes, filename: str = "") -> tuple[list[dict], dict]:
    fmt = detect_format(content, filename)
    if fmt not in ("xlsx", "csv"):
        raise FormatError(_("Загрузите файл Excel (.xlsx) или CSV по шаблону"))
    table = read_table(content, fmt, {})
    if len(table.rows) > MAX_ROWS:
        raise FormatError(_("В файле больше {max} строк — разбейте его на части").format(max=MAX_ROWS))
    rows = []
    for number, values, _extra in table.rows:
        data, errors, warnings = normalize(values)
        if not any(data[f] for f in ("sku", "barcode", "name")):
            continue
        rows.append({"row": number, "data": data, "errors": errors, "warnings": warnings})
    return rows, table.columns


def _lookup(model_qs, field, values):
    if not values:
        return {}
    return dict(model_qs.annotate(_k=Lower(field)).filter(_k__in=values).values_list("_k", "id"))


def plan(rows: list[dict], seller) -> tuple[list[dict], dict]:
    """Решает для каждой строки: создать / обновить / ошибка. Строки меняются на месте."""
    skus = {r["data"]["sku"].lower() for r in rows if r["data"]["sku"]}
    barcodes = {r["data"]["barcode"].lower() for r in rows if r["data"]["barcode"]}
    by_sku = _lookup(ProductVariant.objects.all(), "code", skus)
    by_barcode = _lookup(ProductVariant.objects.all(), "barcode", barcodes)
    extra = ProductBarcode.objects.annotate(_k=Lower("code")).filter(_k__in=barcodes).values_list("_k", "variant_id")
    for code, vid in extra:
        by_barcode.setdefault(code, vid)
    found = set(by_sku.values()) | set(by_barcode.values())
    editable = set(editable_variants(seller).filter(pk__in=found).values_list("id", flat=True))
    categories = {
        name.lower() for name in Category.objects.annotate(_k=Lower("name")).values_list("name", flat=True) if name
    }

    seen_sku, seen_barcode, new_categories = {}, {}, set()
    stats = {"rows": len(rows), "create": 0, "update": 0, "errors": 0, "stock": 0, "warnings": 0}
    for r in rows:
        d, errors, warnings = r["data"], list(r["errors"]), list(r["warnings"])
        sku, barcode = d["sku"].lower(), d["barcode"].lower()
        if sku and sku in seen_sku:
            errors.append(_("Артикул повторяется (строка {row})").format(row=seen_sku[sku]))
        if barcode and barcode in seen_barcode:
            errors.append(_("Штрихкод повторяется (строка {row})").format(row=seen_barcode[barcode]))
        if sku:
            seen_sku.setdefault(sku, r["row"])
        if barcode:
            seen_barcode.setdefault(barcode, r["row"])

        v_sku, v_barcode = by_sku.get(sku), by_barcode.get(barcode)
        variant_id = v_sku or v_barcode
        if v_sku and v_barcode and v_sku != v_barcode:
            errors.append(_("Артикул и штрихкод принадлежат разным товарам"))
        elif variant_id and variant_id not in editable:
            errors.append(_("Товар с этим артикулом/штрихкодом принадлежит другому продавцу"))
        elif not variant_id:
            if not d["name"]:
                errors.append(_("Для нового товара укажите название"))
            if not d["price"]:
                errors.append(_("Для нового товара укажите цену"))

        if d["category"] and d["category"].lower() not in categories and d["category"].lower() not in new_categories:
            new_categories.add(d["category"].lower())
            warnings.append(_("Будет создана категория «{name}»").format(name=d["category"]))
        if d["price"] and d["old_price"] and Decimal(d["old_price"]) <= Decimal(d["price"]):
            warnings.append(_("Старая цена не больше цены — скидка не будет показана"))

        r["errors"], r["warnings"], r["variant_id"] = errors, warnings, variant_id
        r["action"] = "error" if errors else ("update" if variant_id else "create")
        stats["errors" if errors else r["action"]] += 1
        stats["warnings"] += bool(warnings)
        if not errors and d["quantity"] is not None:
            stats["stock"] += 1
    return rows, stats


# ---------------------------------------------------------------------------
# Применение
# ---------------------------------------------------------------------------


class _Cache:
    def __init__(self):
        self.categories, self.brands = {}, {}

    def category(self, name):
        if not name:
            return None
        key = name.lower()
        if key not in self.categories:
            self.categories[key] = Category.objects.filter(name__iexact=name).first() or Category.objects.create(
                name=name
            )
        return self.categories[key]

    def brand(self, name):
        if not name:
            return None
        key = name.lower()
        if key not in self.brands:
            self.brands[key] = Brand.objects.filter(name__iexact=name).first() or Brand.objects.create(name=name)
        return self.brands[key]


def _variant_name(d):
    return f"{d['name']} {d['variant']}".strip() if d["variant"] else d["name"]


def _update(variant: ProductVariant, d: dict, cache: _Cache):
    product = variant.product
    fields = {}
    price = Decimal(d["price"]) if d["price"] else None
    old = Decimal(d["old_price"]) if d["old_price"] else None
    base = price if price is not None else variant.price
    if base is not None and (price is not None or old is not None):
        current_old = variant.original_price if variant.discount else None
        fields.update(discount_fields(Decimal(base).quantize(Decimal("0.01")), old or current_old))
    if d["cost"] is not None:
        fields["cost_price"] = Decimal(d["cost"])
    if d["barcode"] and d["barcode"] != variant.barcode:
        fields["barcode"] = d["barcode"]
    if d["sku"] and (variant.code or "").lower() != d["sku"].lower():
        fields["code"] = d["sku"]
    if d["brand"]:
        fields["brand"] = cache.brand(d["brand"])
    if d["visible"] is not None:
        fields["available"] = d["visible"]
    if d["name"]:
        fields["variant_name"] = _variant_name(d)
    if fields:
        ProductVariant.objects.filter(pk=variant.pk).update(**fields)

    product_fields = {}
    if d["category"]:
        product_fields["category"] = cache.category(d["category"])
    if d["description"]:
        product_fields["description"] = d["description"]
    if d["name"] and product.variants.count() == 1:
        product_fields["name"] = d["name"]
    if product_fields:
        Product.objects.filter(pk=product.pk).update(**product_fields)


def _create(d: dict, seller, publish: bool, product, cache: _Cache) -> ProductVariant:
    if product is None:
        product = Product.objects.create(
            name=d["name"], seller=seller, category=cache.category(d["category"]), description=d["description"] or None
        )
    old = Decimal(d["old_price"]) if d["old_price"] else None
    variant = ProductVariant(
        product=product,
        variant_name=_variant_name(d),
        code=d["sku"] or None,
        barcode=d["barcode"] or None,
        cost_price=Decimal(d["cost"]) if d["cost"] is not None else None,
        brand=cache.brand(d["brand"]),
        available=publish if d["visible"] is None else d["visible"],
        **discount_fields(Decimal(d["price"]).quantize(Decimal("0.01")), old),
    )
    variant.save()
    return variant


def apply(run) -> tuple[dict, list, list]:
    """Применяет строки запуска. → (статистика, строки с результатом, [(variant_id, ссылки на фото)])."""
    rows, _stats = plan(run.rows, run.seller)
    warehouse = None
    if any(r["action"] != "error" and r["data"]["quantity"] is not None for r in rows):
        warehouse = Warehouse.objects.filter(seller=run.seller, pk=run.options.get("warehouse_id")).first()
    publish = bool(run.options.get("publish"))
    cache = _Cache()
    variants = {
        v.pk: v
        for v in ProductVariant.objects.select_related("product").filter(
            pk__in=[r["variant_id"] for r in rows if r["variant_id"]]
        )
    }
    groups = {}
    for r in rows:  # группы, в которых уже есть существующий товар
        if r["action"] == "update" and r["data"]["group"]:
            groups.setdefault(r["data"]["group"].lower(), variants[r["variant_id"]].product)

    stats = {"created": 0, "updated": 0, "failed": 0, "stock": 0}
    photos = []
    for r in rows:
        if r["action"] not in ("create", "update"):
            continue
        d = r["data"]
        try:
            with transaction.atomic():
                if r["action"] == "update":
                    variant = variants[r["variant_id"]]
                    _update(variant, d, cache)
                else:
                    group = d["group"].lower()
                    variant = _create(d, run.seller, publish, groups.get(group) if group else None, cache)
                    if group:
                        groups.setdefault(group, variant.product)
                if d["quantity"] is not None:
                    if warehouse is None:
                        raise stock.StockError(_("Не выбран склад для остатков"))
                    comment = _("Импорт товаров #{run}, строка {row}").format(run=run.pk, row=r["row"])
                    stock.adjust(variant, warehouse, Decimal(d["quantity"]), user=run.created_by, comment=comment)
                    stats["stock"] += 1
        except (stock.StockError, IntegrityError) as exc:
            r["action"], r["result"] = "error", "failed"
            message = str(exc) if isinstance(exc, stock.StockError) else _("Артикул или штрихкод уже занят")
            r["errors"] = [*r["errors"], message]
            stats["failed"] += 1
            continue
        r["result"] = r["action"]
        r["variant_id"] = variant.pk
        stats["created" if r["action"] == "create" else "updated"] += 1
        if d["photos"]:
            photos.append((variant.pk, d["photos"]))
    return stats, rows, photos
