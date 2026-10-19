import csv
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from app.apps.accounts.tests.factories import UserFactory
from app.apps.imports.fields import detect_columns
from app.apps.imports.models import ImportRun
from app.apps.imports.template import COLUMNS, read_rows
from app.apps.products.models import Category, Product, ProductVariant
from app.apps.products.tests.factories import ProductVariantFactory
from app.apps.warehouse.models import StockBalance, StockMovement
from app.apps.warehouse.tests.factories import SellerFactory, WarehouseFactory

pytestmark = pytest.mark.django_db

HEADER = [str(h) for _f, h, _hint in COLUMNS]


def _xlsx(rows, header=HEADER):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append([row.get(f, "") for f, _h, _hint in COLUMNS] if isinstance(row, dict) else row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("products.xlsx", buf.getvalue())


def _upload(client, rows, **extra):
    return client.post(reverse("imports:run-list"), {"file": _xlsx(rows), **extra}, format="multipart")


def test_template_download_roundtrip(client):
    resp = client.get(reverse("imports:template"))
    assert resp.status_code == 200 and resp["Content-Disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(resp.content))
    assert len(wb.sheetnames) == 2
    # заголовки шаблона распознаются как поля импорта
    header = [c.value for c in wb.worksheets[0][1]]
    assert set(detect_columns(header)) == {f for f, _h, _hint in COLUMNS}
    assert read_rows(resp.content, "t.xlsx") == ([], {f: str(h) for f, h, _hint in COLUMNS})

    resp = client.get(reverse("imports:template"), {"type": "csv"})
    assert resp.content.startswith("﻿".encode()) and b";" in resp.content


def test_export_with_products_contains_only_own_goods(client, seller):
    mine = ProductVariantFactory(product__seller=seller, code="MY-1", price=Decimal("100"))
    ProductVariantFactory(product__seller=SellerFactory(), code="ALIEN-1")
    resp = client.get(reverse("imports:template"), {"type": "csv", "products": "1"})
    rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig")), delimiter=";"))
    assert [r[0] for r in rows[1:]] == [mine.code]


def test_preview_classifies_rows(client, seller):
    ProductVariantFactory(product__seller=seller, code="OLD-1", barcode="4820000000100")
    ProductVariantFactory(product__seller=SellerFactory(), code="ALIEN-1")
    rows = [
        {"sku": "NEW-1", "name": "Змішувач", "price": "1200,50", "category": "Нова категорія"},
        {"sku": "OLD-1", "price": 999},
        {"sku": "NEW-2", "name": "Без ціни"},
        {"sku": "NEW-1", "name": "Дубль", "price": 10},
        {"sku": "ALIEN-1", "price": 5},
        {"sku": "NEW-3", "name": "Кривa ціна", "price": "abc"},
        {"sku": "OLD-1X", "barcode": "4820000000100", "price": 7},
    ]
    resp = _upload(client, rows)
    assert resp.status_code == 201, resp.json()
    run = resp.json()
    assert run["status"] == "preview"
    assert run["stats"]["create"] == 1 and run["stats"]["update"] == 2 and run["stats"]["errors"] == 4

    detail = client.get(reverse("imports:run-rows", args=[run["id"]]), {"filter": "errors"}).json()
    messages = {r["row"]: " ".join(r["errors"]) for r in detail["results"]}
    assert "цену" in messages[4]
    assert "повторяется" in messages[5]
    assert "другому продавцу" in messages[6]
    assert "Цена" in messages[7]
    created = client.get(reverse("imports:run-rows", args=[run["id"]]), {"filter": "create"}).json()["results"][0]
    assert any("категория" in w for w in created["warnings"])


def test_apply_creates_updates_groups_and_keeps_empty_cells(client, seller):
    existing = ProductVariantFactory(
        product__seller=seller, code="OLD-1", price=Decimal("100"), cost_price=Decimal("50"), available=True
    )
    rows = [
        {"sku": "OLD-1", "price": "150", "old_price": "200"},
        {
            "sku": "N-1",
            "name": "Ванна",
            "variant": "150 см",
            "group": "G1",
            "price": 9000,
            "brand": "Kolo",
            "category": "Ванни",
        },
        {"sku": "N-2", "name": "Ванна", "variant": "170 см", "group": "G1", "price": 9900, "category": "Ванни"},
        {"sku": "N-3", "name": "Сифон", "price": 300, "visible": "так"},
    ]
    run_id = _upload(client, rows).json()["id"]
    resp = client.post(reverse("imports:run-apply", args=[run_id]))
    assert resp.status_code == 202
    run = ImportRun.objects.get(pk=run_id)
    assert run.status == ImportRun.DONE and run.stats["result"] == {"created": 3, "updated": 1, "failed": 0, "stock": 0}

    existing.refresh_from_db()
    assert (existing.price, existing.original_price, existing.discount) == (Decimal("150.00"), Decimal("200.00"), 25)
    assert existing.cost_price == Decimal("50.0000") and existing.available is True  # пустые ячейки не тронули

    n1, n2, n3 = (ProductVariant.objects.get(code=c) for c in ("N-1", "N-2", "N-3"))
    assert n1.product_id == n2.product_id and n1.product.seller == seller
    assert n1.variant_name == "Ванна 150 см" and n1.brand.name == "Kolo"
    assert n1.product.category == Category.objects.get(name="Ванни")
    assert n1.available is False  # новые по умолчанию скрыты до проверки
    assert n3.available is True  # явно «так»

    resp = client.post(reverse("imports:run-apply", args=[run_id]))
    assert resp.status_code == 400  # повторно применить нельзя


def test_quantity_is_inventory_and_idempotent(client, seller, wh):
    rows = [{"sku": "Q-1", "name": "Кран", "price": 100, "quantity": "7"}]
    first = _upload(client, rows).json()
    assert client.post(reverse("imports:run-apply", args=[first["id"]])).status_code == 400  # нужен склад
    assert client.post(reverse("imports:run-apply", args=[first["id"]]), {"warehouse": wh.pk}).status_code == 202
    variant = ProductVariant.objects.get(code="Q-1")
    assert StockBalance.objects.get(variant=variant, warehouse=wh).quantity == Decimal("7.000")

    again = _upload(client, rows, warehouse=wh.pk).json()
    client.post(reverse("imports:run-apply", args=[again["id"]]))
    variant.refresh_from_db()
    assert variant.stock == Decimal("7.000")  # тот же файл — остаток не удвоился
    assert StockMovement.objects.filter(variant=variant).count() == 1


def test_foreign_warehouse_and_access(client, seller):
    foreign = WarehouseFactory()
    resp = _upload(client, [{"sku": "A", "name": "A", "price": 1}], warehouse=foreign.pk)
    assert resp.status_code == 400

    buyer = APIClient()
    buyer.force_authenticate(UserFactory())
    assert buyer.get(reverse("imports:template")).status_code == 403
    assert buyer.post(reverse("imports:run-list"), {"file": _xlsx([])}, format="multipart").status_code == 403

    run = _upload(client, [{"sku": "A", "name": "A", "price": 1}]).json()
    other = APIClient()
    other.force_authenticate(SellerFactory().user)
    assert other.get(reverse("imports:run-detail", args=[run["id"]])).status_code == 404


def test_bad_files_are_rejected(client):
    bad = SimpleUploadedFile("x.pdf", b"%PDF-1.4")
    assert client.post(reverse("imports:run-list"), {"file": bad}, format="multipart").status_code == 400
    garbage = SimpleUploadedFile("x.xlsx", b"not an excel")
    assert client.post(reverse("imports:run-list"), {"file": garbage}, format="multipart").status_code == 400
    no_header = _xlsx([["a", "b"]], header=["щось", "інше"])
    assert client.post(reverse("imports:run-list"), {"file": no_header}, format="multipart").status_code == 400


def test_cancel_preview_deletes_run(client):
    run = _upload(client, [{"sku": "A", "name": "A", "price": 1}]).json()
    assert client.delete(reverse("imports:run-detail", args=[run["id"]])).status_code == 204
    assert not ImportRun.objects.filter(pk=run["id"]).exists()
    assert Product.objects.count() == 0
